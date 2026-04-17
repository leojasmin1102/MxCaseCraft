"""需求管理模块"""
import json
import openpyxl
from typing import List, Dict


class RequirementManager:
    def __init__(self):
        self.requirements = []

    def extract_requirements_from_excel(self, file_path: str) -> List[Dict]:
        """从Excel文件中按行提取需求

        每一行（除表头外）作为一条需求
        """
        print(f"[需求管理] 从Excel文件提取需求: {file_path}")

        wb = openpyxl.load_workbook(file_path, data_only=True)
        requirements = []
        req_id = 1

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            print(f"[需求管理] 处理工作表: {sheet_name}")

            # 读取表头（第一行）
            headers = []
            for cell in ws[1]:
                headers.append(cell.value if cell.value else "")

            # 从第二行开始读取数据
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                # 跳过空行
                if all(cell is None or str(cell).strip() == "" for cell in row):
                    continue

                # 构建需求内容
                row_data = {}
                content_parts = []

                for col_idx, (header, cell_value) in enumerate(zip(headers, row)):
                    if cell_value is not None and str(cell_value).strip():
                        row_data[header] = str(cell_value).strip()
                        content_parts.append(f"{header}: {str(cell_value).strip()}")

                if content_parts:
                    content = "\n".join(content_parts)
                    requirements.append({
                        'id': req_id,
                        'sheet': sheet_name,
                        'row': row_idx,
                        'content': content,
                        'raw_data': row_data,
                        'confirmed': False
                    })
                    req_id += 1

        print(f"[需求管理] 提取到 {len(requirements)} 条需求")
        self.requirements = requirements
        return requirements

    def extract_requirements_from_content(self, file_content: str) -> List[Dict]:
        """从文件内容中提取需求（旧方法，保留兼容性）

        简单实现：按段落分割，每个非空段落作为一条需求
        """
        print("[需求管理] 开始提取需求...")

        # 按双换行符分割段落
        paragraphs = file_content.split('\n\n')

        requirements = []
        req_id = 1

        for para in paragraphs:
            para = para.strip()
            if para and len(para) > 20:  # 过滤太短的内容
                requirements.append({
                    'id': req_id,
                    'content': para,
                    'confirmed': False
                })
                req_id += 1

        print(f"[需求管理] 提取到 {len(requirements)} 条需求")
        self.requirements = requirements
        return requirements

    def show_requirement(self, req: Dict) -> None:
        """显示单条需求"""
        print("\n" + "="*80)
        print(f"需求 #{req['id']}")
        if 'sheet' in req and 'row' in req:
            print(f"来源: 工作表 [{req['sheet']}] 第 {req['row']} 行")
        print("="*80)
        print(req['content'])
        print("="*80)

    def confirm_requirements_interactive(self) -> List[Dict]:
        """交互式确认需求"""
        confirmed_requirements = []

        print("\n[需求管理] 开始交互式确认需求")
        print("提示：")
        print("  - 输入 'y' 或直接回车：确认当前需求")
        print("  - 输入 'e'：编辑当前需求")
        print("  - 输入 's'：跳过当前需求")
        print("  - 输入 'q'：退出并使用已确认的需求")
        print()

        for req in self.requirements:
            self.show_requirement(req)

            while True:
                choice = input(f"\n请选择操作 [y/e/s/q] (默认: y): ").strip().lower()

                if choice == '' or choice == 'y':
                    # 确认需求
                    req['confirmed'] = True
                    confirmed_requirements.append(req)
                    print(f"✓ 需求 #{req['id']} 已确认")
                    break

                elif choice == 'e':
                    # 编辑需求
                    print("\n请输入修改后的需求内容（输入 'END' 结束编辑）:")
                    lines = []
                    while True:
                        line = input()
                        if line.strip() == 'END':
                            break
                        lines.append(line)

                    new_content = '\n'.join(lines).strip()
                    if new_content:
                        req['content'] = new_content
                        print(f"✓ 需求 #{req['id']} 已更新")
                        self.show_requirement(req)
                    else:
                        print("× 内容为空，保持原内容")

                elif choice == 's':
                    # 跳过需求
                    print(f"⊘ 需求 #{req['id']} 已跳过")
                    break

                elif choice == 'q':
                    # 退出
                    print(f"\n[需求管理] 提前退出，已确认 {len(confirmed_requirements)} 条需求")
                    return confirmed_requirements

                else:
                    print("无效的选择，请重新输入")

        print(f"\n[需求管理] 完成！共确认 {len(confirmed_requirements)} 条需求")
        return confirmed_requirements

    def save_requirements(self, requirements: List[Dict], output_path: str):
        """保存需求到文件"""
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(requirements, f, ensure_ascii=False, indent=2)
        print(f"[需求管理] 需求已保存到 {output_path}")

    def load_requirements(self, input_path: str) -> List[Dict]:
        """从文件加载需求"""
        with open(input_path, 'r', encoding='utf-8') as f:
            requirements = json.load(f)
        print(f"[需求管理] 从 {input_path} 加载了 {len(requirements)} 条需求")
        self.requirements = requirements
        return requirements
