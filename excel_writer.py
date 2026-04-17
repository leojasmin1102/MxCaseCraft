"""Excel输出模块"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime
from typing import List, Dict


class ExcelWriter:
    def __init__(self, template_path: str = None):
        """初始化Excel写入器

        Args:
            template_path: 模板文件路径，如果提供则读取模板格式
        """
        self.template_path = template_path
        self.headers = [
            '用例编号', '所属产品', '相关研发需求', '用例标题',
            '测试步骤', '验收标准', '预期结果', '优先级',
            '用例类型', '执行人', '执行结果', 'BUG记录',
            '测试结果', '由谁创建', '创建日期', '开发对接人员', '算例提供人员'
        ]

        # 如果有模板，读取模板的表头
        if template_path:
            self._load_template_headers()

    def _load_template_headers(self):
        """从模板文件加载表头"""
        try:
            wb = openpyxl.load_workbook(self.template_path)
            ws = wb.active
            self.headers = [cell.value for cell in ws[1] if cell.value]
            print(f"[Excel] 从模板加载了 {len(self.headers)} 个列")
        except Exception as e:
            print(f"[Excel] 加载模板失败: {e}，使用默认表头")

    def write_test_cases(
        self,
        test_cases: List[Dict],
        output_path: str,
        product_name: str = "MxSim.Mechanical",
        requirement: str = "",
        creator: str = "",
        developer: str = ""
    ):
        """将测试用例写入Excel文件

        Args:
            test_cases: 测试用例列表
            output_path: 输出文件路径
            product_name: 产品名称
            requirement: 相关研发需求
            creator: 创建人
            developer: 开发对接人员
        """
        print(f"[Excel] 开始写入Excel文件: {output_path}")

        # 创建工作簿
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "测试用例"

        # 设置表头样式
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # 边框样式
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # 写入表头
        for col_idx, header in enumerate(self.headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        # 设置列宽
        column_widths = {
            'A': 30,  # 用例编号
            'B': 20,  # 所属产品
            'C': 25,  # 相关研发需求
            'D': 40,  # 用例标题
            'E': 50,  # 测试步骤
            'F': 30,  # 验收标准
            'G': 30,  # 预期结果
            'H': 10,  # 优先级
            'I': 12,  # 用例类型
            'J': 12,  # 执行人
            'K': 12,  # 执行结果
            'L': 20,  # BUG记录
            'M': 12,  # 测试结果
            'N': 12,  # 由谁创建
            'O': 15,  # 创建日期
            'P': 12,  # 开发对接人员
            'Q': 12,  # 算例提供人员
        }

        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width

        # 写入测试用例数据
        current_date = datetime.now()

        for idx, test_case in enumerate(test_cases, 1):
            row_idx = idx + 1

            # 生成用例编号（简化版）
            case_id = f"TC-{current_date.strftime('%Y%m%d')}-{idx:03d}"

            # 将steps列表转换为带编号的字符串
            steps_text = ""
            if isinstance(test_case.get('steps'), list):
                steps_text = "\n".join([f"{i}. {step}" for i, step in enumerate(test_case['steps'], 1)])
            else:
                steps_text = test_case.get('steps', '')

            # 写入数据
            row_data = {
                '用例编号': case_id,
                '所属产品': product_name,
                '相关研发需求': requirement,
                '用例标题': test_case.get('title', ''),
                '测试步骤': steps_text,
                '验收标准': test_case.get('expected', ''),
                '预期结果': test_case.get('expected', ''),
                '优先级': '中',
                '用例类型': '功能测试',
                '执行人': '',
                '执行结果': '',
                'BUG记录': '',
                '测试结果': '',
                '由谁创建': creator,
                '创建日期': current_date,
                '开发对接人员': developer,
                '算例提供人员': '/'
            }

            # 写入每一列
            for col_idx, header in enumerate(self.headers, 1):
                value = row_data.get(header, '')
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)

                # 日期格式
                if header == '创建日期' and isinstance(value, datetime):
                    cell.number_format = 'YYYY-MM-DD'

        # 设置行高
        ws.row_dimensions[1].height = 30  # 表头行高
        for row_idx in range(2, len(test_cases) + 2):
            ws.row_dimensions[row_idx].height = 80  # 数据行高

        # 保存文件
        wb.save(output_path)
        print(f"[Excel] 成功写入 {len(test_cases)} 个测试用例到 {output_path}")

        return output_path
