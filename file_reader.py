"""文件读取模块"""
import os
import json
from openpyxl import load_workbook
from docx import Document


class FileReader:
    @staticmethod
    def read_excel(file_path: str) -> str:
        """读取Excel文件内容"""
        print(f"[文件读取] 正在读取Excel: {os.path.basename(file_path)}")
        wb = load_workbook(file_path, data_only=True)
        content = []

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            content.append(f"\n=== 工作表: {sheet_name} ===\n")

            for row in sheet.iter_rows(values_only=True):
                row_data = [str(cell) if cell is not None else "" for cell in row]
                if any(row_data):  # 跳过空行
                    content.append(" | ".join(row_data))

        result = "\n".join(content)
        print(f"[文件读取] Excel读取完成，共 {len(result)} 字符")
        return result

    @staticmethod
    def read_word(file_path: str) -> str:
        """读取Word文件内容"""
        print(f"[文件读取] 正在读取Word: {os.path.basename(file_path)}")
        doc = Document(file_path)
        content = []

        for para in doc.paragraphs:
            if para.text.strip():
                content.append(para.text)

        # 读取表格
        for table in doc.tables:
            content.append("\n=== 表格 ===")
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                content.append(" | ".join(row_data))

        result = "\n".join(content)
        print(f"[文件读取] Word读取完成，共 {len(result)} 字符")
        return result

    @staticmethod
    def read_input_files(input_dir: str = "input") -> dict:
        """读取input目录中的所有文件"""
        print(f"[文件读取] 开始扫描目录: {input_dir}")

        if not os.path.exists(input_dir):
            raise FileNotFoundError(f"目录不存在: {input_dir}")

        files = os.listdir(input_dir)
        print(f"[文件读取] 找到 {len(files)} 个文件")

        content = {
            "excel_files": [],
            "word_files": [],
            "all_content": []
        }

        for file in files:
            file_path = os.path.join(input_dir, file)

            if file.endswith(('.xlsx', '.xls')):
                excel_content = FileReader.read_excel(file_path)
                content["excel_files"].append({"name": file, "content": excel_content})
                content["all_content"].append(f"\n### 文件: {file} ###\n{excel_content}")

            elif file.endswith(('.docx', '.doc')):
                word_content = FileReader.read_word(file_path)
                content["word_files"].append({"name": file, "content": word_content})
                content["all_content"].append(f"\n### 文件: {file} ###\n{word_content}")

        print(f"[文件读取] 读取完成: {len(content['excel_files'])} 个Excel, {len(content['word_files'])} 个Word")
        return content
